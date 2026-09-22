import openpyxl
import glob

f = glob.glob('*.xlsx')[1]
print(f'读取文件：{f}')
wb = openpyxl.load_workbook(f)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} ===')
    print(f'行数：{ws.max_row}, 列数：{ws.max_column}')
    for i in range(1, min(20, ws.max_row+1)):
        row = [ws.cell(row=i, column=j).value for j in range(1, min(25, ws.max_column+1))]
        if any(v is not None for v in row):
            print(f'Row {i}: {row}')
